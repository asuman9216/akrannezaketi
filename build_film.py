from docx import Document
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path

root=Path(__file__).parent
doc=Document()
s=doc.sections[0]
s.top_margin=s.bottom_margin=Cm(1.8)
s.left_margin=s.right_margin=Cm(2)
s.page_width=Cm(21); s.page_height=Cm(29.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,2)
doc.add_heading('Akran nezaketi kısa film planı',0)
p('Umut Derya ve Nida ile dört sahnede okulda nezaket')
p('Bu film, sınıfta başlayan bir konuşmanın kantinde sıraya saygı, koridorda yardımlaşma ve sınıfta arkadaşını gruba katma davranışlarına dönüşmesini anlatır. Her sahne 8 saniyedir. Toplam süre 32 saniyedir; ayrıca açılış veya kapanış sahnesi eklenmez.')
h('Kitapla bağlantısı')
p('MEB Ortaöğretim Akran Nezaketi Etkinlik Kitabı temel alınmıştır. Basılı sayfa 15–16’daki Nezaket İfadelerim, 22–23’teki Dil Değişir Nezaket Gelişir, 24–25’teki Nezaket Etkisi ve 75’teki kantinde alışveriş hakkı örneği kurguya yön verir. PDF sayfa numaraları bu numaralardan bir fazladır. Diyaloglar bu film için özgün olarak yazılmıştır.')
h('Karakterlerin değişmeyen görünümü')
p('Üçü de yaklaşık 15 yaşında lise öğrencisidir. Umut kısa koyu kahverengi saçlı, lacivert sweatshirt ve koyu pantolonlu bir erkek öğrencidir. Derya omuz hizasında düz kahverengi saçlı, açık mavi sweatshirt ve koyu pantolonlu bir kız öğrencidir. Nida koyu kıvırcık saçlarını at kuyruğu toplamış, bordo sweatshirt ve koyu pantolonlu bir kız öğrencidir. Umut’un yeşil, Derya’nın mavi, Nida’nın sarı defteri vardır. Yüzler, saçlar ve kıyafetler bütün sahnelerde aynı kalır.')
h('Ortak görsel ve ses dili')
p('16:9 yatay kadraj, gerçekçi sinematik görüntü, doğal gün ışığı, aynı Türkiye lisesi ve aynı okul günü. Kamera öğrencilerin göz hizasında; oyunculuk sade ve doğal. Karakterleri önce tek bir referans görselinde belirleyin ve dört görsel üretiminde aynı referansı kullanın. Her sahnenin görseli, hareketin başlangıç karesini göstermelidir.')
p('1 ve 2. sahnelerde kısa Türkçe diyalog vardır. 3 ve 4. sahneler diyalogsuzdur; yalnızca hafif ortam sesleri duyulur. Konuşmalar üst üste binmez. Video aracı ses üretemiyorsa aynı replikler belirtilen zamanlara sonradan eklenir. Müzik ve dış ses kullanılmaz.')
p('Anlatım, zorbalıkla ilişkili incitici söz ve dışlama karşısında destekleyici davranışlara odaklanır. Nezaket, zarar gören öğrencinin tek başına çözüm üretme yükümlülüğü gibi sunulmaz; ilk sahnede bir yetişkinden birlikte destek isteme yer alır.')
scenes=[
('1 Sınıfta yanında olmak','00:00–00:08 | 8 saniye | Diyaloglu',
'Derya, tekrarlanan alaydan rahatsız olduğunu Umut’a söyler. Umut dinler ve birlikte destek istemeyi önerir. Nida bu sahnede görünmez.',
['0–1 sn: Derya sırasındaki mavi deftere bakar, sonra Umut’a döner.','1–4 sn | Derya: “Benimle yine alay ettiler. Üzüldüm.”','4–7 sn | Umut: “Yanındayım. Öğretmenimizle birlikte konuşalım.”','7–8 sn: Derya başıyla onaylar; Umut destekleyici bir ifadeyle yanında kalır.'],
'16:9 yatay, gerçekçi sinematik fotoğraf. Türkiye’de bir lise sınıfı, pencereden yumuşak gün ışığı, ahşap sıralar. Omuz hizasında düz kahverengi saçlı, açık mavi sweatshirtlü 15 yaşındaki Derya ön sırada oturuyor; önünde kapalı mavi defteri var. Yanında kısa koyu kahverengi saçlı, lacivert sweatshirtlü 15 yaşındaki Umut oturuyor; önünde yeşil defteri var. Derya hafif üzgün, konuşmaya başlamadan hemen önce Umut’a dönüyor. Umut dikkatle ona bakıyor. Belden yukarı ikili orta plan, ikisinin yüzü net. Arka plan hafif flu. Abartılı ağlama, yazı, altyazı, logo ve konuşma balonu yok.',
'Referans görseli başlangıç karesi olarak kullan. Tam 8 saniye, 16:9, tek kesintisiz ikili orta plan; kamera çok hafif yaklaşsın. 0–1 saniyede Derya Umut’a dönsün. 1–4 saniyede yalnızca Derya doğal ve hafif üzgün sesle “Benimle yine alay ettiler. Üzüldüm.” desin. 4–7 saniyede yalnızca Umut sakin ve destekleyici sesle “Yanındayım. Öğretmenimizle birlikte konuşalım.” desin. 7–8 saniyede Derya başıyla onaylasın. Türkçe dudak eşleşmesi; söz sırası dışında ağızlar kapalı. Hafif sınıf ortam sesi. Yüzler ve kıyafetler sabit kalsın. Dış ses, müzik, yazı ve sahne geçişi ekleme.'),
('2 Kantinde sıraya saygı','00:08–00:16 | 8 saniye | Diyaloglu',
'Kantinde Derya, Nida’nın arkasında bekler. Umut, Derya’yı görünce yanına ilerlemek ister; uyarıyı kabul ederek sıranın sonuna geçer.',
['0–1 sn: Nida önde, Derya arkasındadır. Umut sıranın yanından Derya’ya yaklaşır.','1–4 sn | Derya: “Umut, sıranın sonu burada. Birlikte bekleyelim.”','4–6 sn | Umut: “Haklısın, özür dilerim.”','6–8 sn: Umut Derya’nın arkasına geçer. Nida hafifçe gülümser.'],
'16:9 yatay, gerçekçi sinematik fotoğraf. Aynı lisenin aydınlık kantini. Tezgâh kadrajın solunda; öğrenci sırasının yönü açıkça anlaşılıyor. Koyu kıvırcık saçları at kuyruğu toplanmış bordo sweatshirtlü Nida önde, açık mavi sweatshirtlü Derya onun arkasında. Lacivert sweatshirtlü Umut sıraya henüz katılmamış, Derya’nın sağ yanında bir adım uzakta duruyor. Derya’nın arkasındaki sıra sonu kadrajda görünür. Üç karakter aynı yaşta ve önceki referansla aynı yüzlere sahip. Göz hizasında orta geniş plan, sakin okul ortamı. Yazı, altyazı, logo yok.',
'Bu görseli başlangıç karesi olarak kullan. Tam 8 saniye, tek kesintisiz 16:9 orta geniş plan. 0–1 saniyede Umut Derya’ya doğru küçük bir adım atsın. 1–4 saniyede Derya sakin sesle “Umut, sıranın sonu burada. Birlikte bekleyelim.” desin ve eliyle arkasındaki boş yeri göstersin. 4–6 saniyede Umut “Haklısın, özür dilerim.” desin. 6–8 saniyede Umut Derya’nın arkasındaki sıra sonuna geçsin; Nida hafifçe gülümsesin. Son sıra düzeni tezgâhtan geriye Nida, Derya, Umut olsun. Kamera sabit. Hafif kantin uğultusu; başka anlaşılır konuşma olmasın. Doğal Türkçe dudak eşleşmesi. İtme, abartılı tepki, dış ses, müzik ve yazı ekleme.'),
('3 Koridorda yardım etmek','00:16–00:24 | 8 saniye | Diyalogsuz',
'Sınıfa dönerken Nida’nın sarı defteri elinden düşer. Derya defteri alıp geri verir; Umut da yanlarında durup bekler. Küçük bir yardım davranışı sözcük kullanmadan anlatılır.',
['0–2 sn: Nida’nın sarı defteri elinden kayıp yere düşer.','2–5 sn: Derya eğilip defteri alır; Umut durur ve geçiş için alan bırakır.','5–7 sn: Derya defteri Nida’ya uzatır. Nida gülümseyip başıyla teşekkür eder.','7–8 sn: Üçü sınıfa doğru yönelir.'],
'16:9 yatay, gerçekçi sinematik fotoğraf. Aynı lisenin gün ışığı alan koridoru, sağda sınıf kapısı, solda pencereler. Bordo sweatshirtlü Nida ön planda yürümek üzere; sarı defterini parmak uçlarında gevşekçe tutuyor, defter henüz düşmemiş. Açık mavi sweatshirtlü Derya Nida’nın yanında, lacivert sweatshirtlü Umut yarım adım geride. Üçünün de yüzleri referansla aynı. Tam boy orta geniş plan; defterin düşeceği zemin ve eller açıkça görünür. Doğal ifadeler, sade arka plan. Yazı, altyazı ve logo yok.',
'Başlangıç görselinden tam 8 saniyelik, tek kesintisiz 16:9 video üret. Kamera sabit orta geniş planda kalsın. 0–2 saniyede Nida’nın elindeki tek sarı defter yere düşsün. 2–5 saniyede Derya doğal biçimde eğilip defteri yerden alsın; Umut durup bir adım yana çekilerek alan bıraksın. 5–7 saniyede Derya defteri Nida’ya geri versin; Nida gülümseyip başıyla teşekkür etsin. 7–8 saniyede üçü sağdaki sınıf kapısına yönelsin. Konuşma, dudak hareketiyle konuşma, dış ses ve müzik olmasın. Yalnızca hafif ayak sesi ve defterin yere değme sesi. Defter çoğalmasın, el değiştirmesi kesintisiz ve fiziksel olarak tutarlı olsun. Yazı ekleme.'),
('4 Sınıfta herkese yer açmak','00:24–00:32 | 8 saniye | Diyalogsuz',
'Sınıfa döndüklerinde Umut ve Nida aynı sıradadır. Derya masanın yanında durur. Nida boş sandalyeyi dışarı çekip Derya’yı davet eder; üçü aynı çalışmaya katılır. Başlangıçta destek alan Derya, kendisine yer açılan grupta rahatlar.',
['0–2 sn: Nida masanın yanındaki boş sandalyeyi Derya’ya doğru hafifçe çeker.','2–4 sn: Nida açık avuç hareketi ve gülümsemeyle Derya’yı davet eder.','4–6 sn: Derya sandalyeye oturur. Umut açık yeşil defteri üçünün görebileceği yere kaydırır.','6–8 sn: Üçü birlikte deftere bakar; Derya rahatça gülümser.'],
'16:9 yatay, gerçekçi sinematik fotoğraf. İlk sahnedeki aynı lise sınıfı ve aynı gün ışığı. Bir çalışma masasında lacivert sweatshirtlü Umut ve bordo sweatshirtlü Nida oturuyor. Açık mavi sweatshirtlü Derya masanın yanındaki boş sandalyenin arkasında ayakta duruyor. Nida’nın eli sandalyenin sırtlığına yakın, henüz çekmemiş. Masada Umut’un açık yeşil defteri, kapalı mavi ve sarı defterler var; defterlerde okunabilir yazı görünmüyor. Üç yüz, saç ve kıyafet önceki sahnelerle aynı. Boş sandalye ve Derya’nın oturacağı alanı gösteren göz hizasında orta geniş plan. Yazı, altyazı ve logo yok.',
'Bu görseli başlangıç karesi olarak kullan. Tam 8 saniye, tek kesintisiz 16:9 orta geniş plan; kamera sabit kalsın. 0–2 saniyede Nida boş sandalyeyi Derya’ya doğru hafifçe çeksin. 2–4 saniyede gülümseyerek açık avucuyla oturmasını işaret etsin. 4–6 saniyede Derya sandalyeye otursun; Umut açık yeşil defterini masanın ortasına kaydırsın. 6–8 saniyede üçü aynı deftere baksın ve Derya rahatça gülümsesin. Hiç konuşma, dış ses veya müzik olmasın; hafif sandalye ve sayfa sesi yeterli. Yüzleri, kıyafetleri, el ve sandalye geometrisini koru. Ek karakter, yazı, altyazı ve kapanış kartı ekleme. Son kareyi sekizinci saniyede bitir.')
]
for title,meta,story,beats,img,vid in scenes:
 doc.add_page_break(); doc.add_heading(title,1); p(meta)
 h('Olay ve amaç'); p(story)
 h('Saniye saniye akış')
 for b in beats: p(b)
 h('Görsel oluşturma istemi'); p(img)
 h('Videolaştırma istemi'); p(vid)
doc.save(root/'Akran_Nezaketi_Dort_Sahne.docx')
print(root/'Akran_Nezaketi_Dort_Sahne.docx')
